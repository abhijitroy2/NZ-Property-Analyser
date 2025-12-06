"""
Generate Executive Summary Excel file from multiple analysis output files
Collates all 'Good Deals' from flip and rental modes into separate tabs
"""
import pandas as pd
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime
from typing import List, Tuple, Optional

# Color definitions
GREEN_FILL = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def read_good_deals_from_file(file_path: str, mode: str) -> Optional[pd.DataFrame]:
    """
    Read good deals from an Excel file.
    
    Args:
        file_path: Path to Excel file
        mode: 'flip' or 'rental'
        
    Returns:
        DataFrame with good deals, or None if error
    """
    try:
        # Read the Summary sheet first
        if not os.path.exists(file_path):
            return None
            
        df = pd.read_excel(file_path, sheet_name='Summary')
        
        # Filter for good deals based on mode
        if mode == 'flip':
            if 'Is_Good_Deal' in df.columns:
                good_deals = df[df['Is_Good_Deal'] == True].copy()
            else:
                return None
        else:  # rental
            if 'Is_Successful' in df.columns:
                good_deals = df[df['Is_Successful'] == True].copy()
            else:
                return None
        
        # Add source file column
        good_deals['Source_File'] = os.path.basename(file_path)
        
        return good_deals if not good_deals.empty else None
        
    except Exception as e:
        print(f"Error reading {file_path}: {e}")
        return None


def write_dataframe_to_sheet(ws, df: pd.DataFrame):
    """
    Write DataFrame to worksheet with formatting.
    
    Args:
        ws: Worksheet object
        df: DataFrame to write
    """
    # Convert lists and other non-serializable types to strings
    df_clean = df.copy()
    for col in df_clean.columns:
        df_clean[col] = df_clean[col].apply(
            lambda x: ', '.join(str(v) for v in x) if isinstance(x, list) else x
        )
    
    # Write data
    for r_idx, row in enumerate(dataframe_to_rows(df_clean, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            if isinstance(value, (list, dict)):
                cell.value = str(value)
            else:
                cell.value = value
            
            # Format header row
            if r_idx == 1:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                # Highlight all data rows in green (they're all good deals)
                cell.fill = GREEN_FILL
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width


def generate_executive_summary(output_files: List[str], output_path: str) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """
    Generate Executive Summary Excel file from multiple output files.
    
    Args:
        output_files: List of paths to analysis output Excel files
        output_path: Path where Executive Summary should be saved
        
    Returns:
        Tuple of (flip_good_deals_df, rental_good_deals_df)
    """
    flip_good_deals = []
    rental_good_deals = []
    
    # Process each output file
    for file_path in output_files:
        # Determine mode from filename
        filename = os.path.basename(file_path).lower()
        
        if '_flip_' in filename or filename.endswith('_flip.xlsx'):
            # Flip mode file
            deals = read_good_deals_from_file(file_path, 'flip')
            if deals is not None:
                flip_good_deals.append(deals)
        elif '_rental_' in filename or filename.endswith('_rental.xlsx'):
            # Rental mode file
            deals = read_good_deals_from_file(file_path, 'rental')
            if deals is not None:
                rental_good_deals.append(deals)
    
    # Combine all flip deals
    flip_df = pd.concat(flip_good_deals, ignore_index=True) if flip_good_deals else pd.DataFrame()
    
    # Combine all rental deals
    rental_df = pd.concat(rental_good_deals, ignore_index=True) if rental_good_deals else pd.DataFrame()
    
    # Create Excel workbook
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet
    
    # Create Flip tab
    if not flip_df.empty:
        flip_sheet = wb.create_sheet("Good Deals - Flip")
        write_dataframe_to_sheet(flip_sheet, flip_df)
    
    # Create Rental tab
    if not rental_df.empty:
        rental_sheet = wb.create_sheet("Good Deals - Rental")
        write_dataframe_to_sheet(rental_sheet, rental_df)
    
    # Create Metadata sheet
    metadata_sheet = wb.create_sheet("Metadata")
    metadata = [
        ["Executive Summary - NZ Property Analyser"],
        [""],
        ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["Source Files Processed", len(output_files)],
        ["Flip Good Deals Found", len(flip_df)],
        ["Rental Good Deals Found", len(rental_df)],
    ]
    
    for r_idx, row in enumerate(metadata, 1):
        for c_idx, value in enumerate(row, 1):
            cell = metadata_sheet.cell(row=r_idx, column=c_idx)
            cell.value = value
            if r_idx == 1:
                cell.font = Font(bold=True, size=14)
    
    metadata_sheet.column_dimensions['A'].width = 30
    metadata_sheet.column_dimensions['B'].width = 40
    
    # Save workbook
    wb.save(output_path)
    print(f"Executive Summary saved to: {output_path}")
    print(f"  - Flip Good Deals: {len(flip_df)}")
    print(f"  - Rental Good Deals: {len(rental_df)}")
    
    return flip_df, rental_df


def extract_good_deals_info(df: pd.DataFrame, mode: str) -> List[dict]:
    """
    Extract key information from good deals DataFrame for email body.
    
    Args:
        df: DataFrame with good deals
        mode: 'flip' or 'rental'
        
    Returns:
        List of dictionaries with property info
    """
    deals_info = []
    
    if df.empty:
        return deals_info
    
    for _, row in df.iterrows():
        deal_info = {
            'address': row.get('Property Address', 'N/A'),
            'url': row.get('URL', 'N/A'),
        }
        
        if mode == 'flip':
            deal_info['purchase_price'] = row.get('Potential_Purchase_Price', 'N/A')
            deal_info['sale_price'] = row.get('Potential_Sale_Price', 'N/A')
            deal_info['profit'] = row.get('Profit', 'N/A')
            deal_info['roi'] = row.get('ROI_Percent', 'N/A')
            deal_info['cash_on_cash_roi'] = row.get('Cash_on_Cash_ROI_Percent', 'N/A')
        else:  # rental
            deal_info['purchase_price'] = row.get('Potential_Purchase_Price', 'N/A')
            deal_info['weekly_rent'] = row.get('Weekly_Rent', 'N/A')
            deal_info['net_yield'] = row.get('Net_Yield_Percent', 'N/A')
            deal_info['net_income'] = row.get('Net_Income_Loss', 'N/A')
        
        deals_info.append(deal_info)
    
    return deals_info


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python generate_executive_summary.py <output_folder> <output_file>")
        print("Example: python generate_executive_summary.py output Executive_Summary.xlsx")
        sys.exit(1)
    
    output_folder = sys.argv[1]
    output_file = sys.argv[2]
    
    # Get all Excel files in output folder
    excel_files = []
    if os.path.exists(output_folder):
        for file in os.listdir(output_folder):
            if file.endswith('.xlsx') and ('_flip_' in file.lower() or '_rental_' in file.lower()):
                excel_files.append(os.path.join(output_folder, file))
    
    if not excel_files:
        print(f"No analysis files found in {output_folder}")
        sys.exit(1)
    
    # Generate executive summary
    output_path = os.path.join(output_folder, output_file)
    flip_df, rental_df = generate_executive_summary(excel_files, output_path)
    
    # Print summary
    print(f"\nSummary:")
    print(f"  Processed {len(excel_files)} analysis files")
    print(f"  Found {len(flip_df)} flip good deals")
    print(f"  Found {len(rental_df)} rental good deals")

